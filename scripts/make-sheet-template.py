"""Regenerates docs/kozy-sheet-template.xlsx from the dev seed src/data/snapshot.json.

    python3 scripts/make-sheet-template.py

Run it whenever a column is added or renamed so the template the client fills in
never drifts from what the site actually parses. Deliberately formula-free: the
file is meant to be edited on a phone, and live counts already live on the site
and on /purge.
"""

import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DATA = json.load(open(ROOT / 'src' / 'data' / 'snapshot.json'))

INK = '1C1917'
GREEN = '1F6F4D'
GREEN_SOFT = 'E7F3EC'
MUTED = '78716C'
LINE = 'E3E0DB'
BAND = 'FAF9F7'
KOSONG_BG, KOSONG_FG = 'DCF3E4', '156F45'
BOOK_BG, BOOK_FG = 'FBEED5', '8A5A16'
ISI_BG, ISI_FG = 'F1EFEC', '57534E'

F = 'Arial'
thin = Side(style='thin', color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()


def header_row(ws, headers, notes, widths, row=1):
    for col, (name, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.font = Font(name=F, size=10, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=GREEN)
        cell.alignment = Alignment(horizontal='left', vertical='center')
        cell.border = box
        if notes.get(name):
            cell.comment = Comment(notes[name], 'Panduan Kozy', width=280, height=110)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 26
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def body(ws, rows, headers, wrap_cols=()):
    for r, record in enumerate(rows, start=2):
        for c, name in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=c, value=record.get(name, ''))
            cell.font = Font(name=F, size=10, color=INK)
            cell.border = box
            cell.alignment = Alignment(
                vertical='top', wrap_text=name in wrap_cols, horizontal='left'
            )
            if r % 2 == 0:
                cell.fill = PatternFill('solid', fgColor=BAND)
        ws.row_dimensions[r].height = 30


# ---------------------------------------------------------------- config
ws = wb.active
ws.title = 'config'
cfg_headers = ['key', 'value']
header_row(
    ws,
    cfg_headers,
    {
        'key': 'Jangan diubah. Empat baris ini sudah pas.',
        'value': 'Isi di kolom ini. wa_default ditulis 628..., tanpa tanda + dan tanpa spasi.',
    },
    [22, 96],
)
body(ws, DATA['config'], cfg_headers, wrap_cols={'value'})

# ---------------------------------------------------------------- lokasi
ws = wb.create_sheet('lokasi')
lok_headers = [
    'slug', 'nama', 'alamat', 'gmaps_url', 'lat', 'lng', 'deskripsi',
    'fasilitas', 'nomor_wa', 'foto_urls', 'aktif',
]
header_row(
    ws,
    lok_headers,
    {
        'slug': 'Alamat halaman. Huruf kecil, tanpa spasi, pakai tanda hubung.\nContoh: batam-centre',
        'gmaps_url': 'Tautan Google Maps untuk tombol rute di halaman lokasi.',
        'lat': 'Koordinat untuk peta. Klik kanan titik di Google Maps untuk menyalinnya.',
        'fasilitas': 'Fasilitas bersama, dipisah koma.',
        'nomor_wa': 'Boleh dikosongkan. Kalau kosong, website memakai wa_default dari tab config.',
        'foto_urls': 'Dipisah koma. Foto pertama jadi foto utama. Link Google Drive harus di-share Anyone with the link.',
        'aktif': 'TRUE tampil di website. FALSE disembunyikan tanpa menghapus datanya.',
    },
    [16, 24, 44, 34, 10, 10, 62, 46, 16, 52, 9],
)
body(ws, DATA['lokasi'], lok_headers, wrap_cols={'alamat', 'deskripsi', 'fasilitas', 'foto_urls'})

aktif_dv = DataValidation(type='list', formula1='"TRUE,FALSE"', allow_blank=False)
ws.add_data_validation(aktif_dv)
aktif_dv.add('K2:K200')

# ---------------------------------------------------------------- kamar
ws = wb.create_sheet('kamar')
kmr_headers = [
    'kode', 'lokasi_slug', 'lantai', 'luas_m2', 'tipe', 'harga_bulanan',
    'status', 'fasilitas', 'foto_urls', 'catatan', 'aktif',
]
header_row(
    ws,
    kmr_headers,
    {
        'kode': 'Kode kamar. Inilah yang tampil di kotak denah kamar.',
        'lokasi_slug': 'Pilih dari dropdown. Harus sama persis dengan slug di tab lokasi.',
        'lantai': 'Angka. Menentukan pengelompokan baris di denah kamar.',
        'tipe': 'Nama tipe bebas, tapi tulis SAMA PERSIS untuk kamar yang setipe.\nSatu tipe sebaiknya satu harga, satu luas, satu daftar fasilitas.',
        'harga_bulanan': 'Angka saja. Boleh 1850000 atau Rp 1.850.000, keduanya terbaca.',
        'status': 'Pilih dari dropdown. Ini kolom yang paling sering diubah.',
        'fasilitas': 'Dipisah koma. Yang tampil di kartu tipe hanya fasilitas yang dimiliki SEMUA kamar bertipe itu.',
        'catatan': 'Satu kalimat, boleh kosong. Contoh: Ditahan sampai 18 September.',
        'aktif': 'TRUE tampil di website. FALSE disembunyikan tanpa menghapus datanya.',
    },
    [10, 17, 9, 10, 16, 16, 14, 54, 46, 34, 9],
)
body(ws, DATA['kamar'], kmr_headers, wrap_cols={'fasilitas', 'foto_urls', 'catatan'})

status_dv = DataValidation(type='list', formula1='"kosong,dibooking,terisi"', allow_blank=False)
ws.add_data_validation(status_dv)
status_dv.add('G2:G400')

slug_dv = DataValidation(type='list', formula1='lokasi!$A$2:$A$50', allow_blank=False)
ws.add_data_validation(slug_dv)
slug_dv.add('B2:B400')

tipe_names = sorted({row['tipe'] for row in DATA['kamar']})
tipe_dv = DataValidation(type='list', formula1=f'"{",".join(tipe_names)}"', allow_blank=False)
ws.add_data_validation(tipe_dv)
tipe_dv.add('E2:E400')

aktif_dv2 = DataValidation(type='list', formula1='"TRUE,FALSE"', allow_blank=False)
ws.add_data_validation(aktif_dv2)
aktif_dv2.add('K2:K400')

for value, bg, fg in (
    ('kosong', KOSONG_BG, KOSONG_FG),
    ('dibooking', BOOK_BG, BOOK_FG),
    ('terisi', ISI_BG, ISI_FG),
):
    ws.conditional_formatting.add(
        'G2:G400',
        CellIsRule(
            operator='equal',
            formula=[f'"{value}"'],
            fill=PatternFill('solid', bgColor=bg),
            font=Font(bold=True, color=fg),
        ),
    )

# ---------------------------------------------------------------- panduan
ws = wb.create_sheet('Panduan', 0)
ws.sheet_view.showGridLines = False
for col, width in zip('ABCD', (4, 30, 78, 4)):
    ws.column_dimensions[col].width = width


def put(row, col, text, *, size=10, bold=False, color=INK, fill=None, wrap=False, height=None):
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = Font(name=F, size=size, bold=bold, color=color)
    cell.alignment = Alignment(vertical='top', wrap_text=wrap)
    if fill:
        cell.fill = PatternFill('solid', fgColor=fill)
    if height:
        ws.row_dimensions[row].height = height
    return cell


def section(row, title):
    for col in (2, 3):
        put(row, col, title if col == 2 else '', size=11, bold=True, color=GREEN, fill=GREEN_SOFT)
    ws.row_dimensions[row].height = 22
    return row + 1


def step(row, label, text):
    put(row, 2, label, size=10, bold=True)
    put(row, 3, text, size=10, color=MUTED, wrap=True, height=30)
    return row + 1


put(2, 2, 'Panduan mengisi Sheet', size=18, bold=True, height=30)
put(
    3, 2,
    'Website membaca file ini langsung. Tidak ada aplikasi lain dan tidak perlu menghubungi developer.',
    size=10, color=MUTED,
)
ws.merge_cells('B3:C3')

r = 5
r = section(r, 'Tiga tab yang dibaca website')
for name, text in (
    ('config', 'Nama properti, kalimat perkenalan, nomor WhatsApp utama, dan alamat kantor.'),
    ('lokasi', 'Satu baris untuk satu bangunan kos.'),
    ('kamar', 'Satu baris untuk satu kamar fisik. Baris inilah yang jadi kotak di denah kamar.'),
):
    r = step(r, name, text)
r += 1

r = section(r, 'Mengubah status kamar')
for label, text in (
    ('1', 'Buka tab kamar, cari baris kamarnya lewat kolom kode.'),
    ('2', 'Ketuk kolom status, pilih kosong, dibooking, atau terisi dari dropdown.'),
    ('3', 'Buka bookmark Perbarui website. Tanpa langkah ini pun perubahan muncul sendiri dalam 5 menit.'),
):
    r = step(r, label, text)
r += 1

r = section(r, 'Menambah tipe kamar')
for label, text in (
    ('Tidak ada tab tipe', 'Tipe kamar terbentuk sendiri dari kolom tipe di tab kamar.'),
    ('Tulis sama persis', 'Deluxe AC dan Deluxe A/C dianggap dua tipe berbeda. Pakai dropdown supaya seragam.'),
    ('Samakan isinya', 'Harga, luas, dan fasilitas semua kamar bertipe sama sebaiknya diisi sama.'),
    ('Fasilitas dipotong', 'Kartu tipe hanya menampilkan fasilitas yang dimiliki SEMUA kamar bertipe itu.'),
)   :
    r = step(r, label, text)
r += 1

r = section(r, 'Menambah lokasi')
for label, text in (
    ('1', 'Tab lokasi, tambah baris. Isi minimal slug, nama, alamat, dan aktif = TRUE.'),
    ('2', 'Kolom nomor_wa boleh dikosongkan kalau lokasi ini memakai nomor utama di tab config.'),
    ('3', 'Tab kamar, tambah baris kamarnya. Kolom lokasi_slug pilih dari dropdown.'),
    ('4', 'Buka bookmark Perbarui website. Halaman /locations/{slug} langsung ada.'),
):
    r = step(r, label, text)
r += 1

r = section(r, 'Yang tidak boleh dilakukan')
for label, text in (
    ('Jangan hapus baris', 'Untuk menyembunyikan lokasi atau kamar, ubah kolom aktif menjadi FALSE.'),
    ('Jangan ubah baris 1', 'Nama kolom di baris pertama dipakai website untuk mengenali datanya.'),
    ('Jangan ubah nama tab', 'Tiga tab harus tetap bernama config, lokasi, dan kamar, huruf kecil.'),
    ('Jangan cabut sharing', 'File harus tetap Anyone with the link, Viewer. Kalau dicabut, website memakai data cadangan.'),
):
    r = step(r, label, text)
r += 1

r = section(r, 'Melihat angka terbaru')
for label, text in (
    ('Website', 'Halaman depan menampilkan jumlah lokasi, total kamar kosong, dan harga terendah.'),
    ('Halaman lokasi', 'Denah kamar menampilkan status tiap kamar, dikelompokkan per lantai.'),
):
    r = step(r, label, text)
r += 1

r = section(r, 'Kalau ada yang aneh')
r = step(
    r, 'Buka /purge',
    'Halaman itu menyebut sumber data yang sedang dipakai, jumlah baris terbaca, dan baris mana yang dilewati karena datanya belum lengkap.',
)
r = step(
    r, 'Website tidak pernah kosong',
    'Kalau Sheet bermasalah, yang tampil adalah data terakhir yang tersimpan di kode.',
)

out = ROOT / 'docs' / 'kozy-sheet-template.xlsx'
wb.save(out)
print('saved', out)
